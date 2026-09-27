import io


def extract_text(filename: str, content: bytes) -> str:
    """Extract plain text from an uploaded PDF, DOCX, or TXT file."""
    name = filename.lower()
    if name.endswith(".pdf"):
        return _extract_pdf(content)
    if name.endswith(".docx"):
        return _extract_docx(content)
    # fall back: treat as plain text
    return content.decode("utf-8", errors="ignore")


def _extract_pdf(content: bytes) -> str:
    import pdfplumber

    text_parts = []
    with pdfplumber.open(io.BytesIO(content)) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text() or ""
            text_parts.append(page_text)
    text = "\n\n".join(text_parts).strip()
    if len(text) < 20:
        # Likely a scanned/image-based PDF with no extractable text layer.
        raise ValueError(
            "This PDF appears to be scanned or image-based (no extractable text layer). "
            "OCR is not wired into this MVP yet — please paste the resume text manually."
        )
    return text


def _extract_docx(content: bytes) -> str:
    import docx

    doc = docx.Document(io.BytesIO(content))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text for c in row.cells))
    return "\n".join(parts).strip()
